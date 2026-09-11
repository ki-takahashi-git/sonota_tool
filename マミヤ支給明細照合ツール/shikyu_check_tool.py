"""
支給部材PDF ⇔ 移動明細Excel 照合ツール
=====================================

「タカハタ電子御中 支給部材_YYMMDD.pdf」(移動伝票, スキャンPDF) と
「タカハタ電子御中 移動明細_YY.M~.xlsx」(移動明細, 指定シート) の内容を
自動で読み取り、品番・数量が一致しているかを照合するツールです。

PDFは画像スキャンのため、Windows標準のOCR機能(日本語言語パック)を
使って文字を読み取ります。OCRは完璧ではないため、読み取りに自信が
持てない行は「要確認」として黄色でハイライトし、最終判断は必ず
目視で確認してください。

必要なライブラリ:
    pip install pymupdf openpyxl pillow winocr

使い方:
    python shikyu_check_tool.py
"""

import asyncio
import difflib
import re
import sys
import threading
import traceback
from dataclasses import dataclass, field
from pathlib import Path
from tkinter import (
    Tk, StringVar, IntVar, Toplevel, filedialog, messagebox, ttk, BOTH, X, Y,
    LEFT, RIGHT, TOP, BOTTOM, END, N, S, E, W, HORIZONTAL, VERTICAL, WORD
)
import tkinter as tk

import fitz  # PyMuPDF
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment
from openpyxl.utils import get_column_letter
from PIL import Image

try:
    import winocr
except ImportError:
    winocr = None

import json
import os


# ---------------------------------------------------------------------------
# バージョン表示・使い方(同梱データ VERSION / 使い方.md)
# ---------------------------------------------------------------------------

HERE = os.path.dirname(os.path.abspath(__file__))


def _resource_path(filename):
    """
    同梱データ(VERSION・使い方.md)のパスを返す。
    PyInstaller(--onefile)実行時は一時展開先(sys._MEIPASS)に、
    通常のPython実行時はこのファイルと同じフォルダに置かれる。
    """
    base = getattr(sys, '_MEIPASS', HERE)
    return os.path.join(base, filename)


def get_version():
    """VERSIONファイルの内容(例: "1.0.0")を返す。読めない場合は "?"。"""
    try:
        with open(_resource_path('VERSION'), encoding='utf-8') as f:
            return f.read().strip()
    except OSError:
        return '?'


# ---------------------------------------------------------------------------
# 設定の保存(前回選択したフォルダを記憶する)
# ---------------------------------------------------------------------------

def _config_path():
    base = os.environ.get('APPDATA') or str(Path.home())
    d = Path(base) / '支給明細照合ツール'
    d.mkdir(parents=True, exist_ok=True)
    return d / 'config.json'


def load_config():
    try:
        with open(_config_path(), 'r', encoding='utf-8') as f:
            return json.load(f)
    except (OSError, ValueError):
        return {}


def save_config(cfg):
    try:
        with open(_config_path(), 'w', encoding='utf-8') as f:
            json.dump(cfg, f, ensure_ascii=False, indent=2)
    except OSError:
        pass


# ---------------------------------------------------------------------------
# 共通ユーティリティ
# ---------------------------------------------------------------------------

DASH_CHARS = str.maketrans({c: '-' for c in 'ー一−‐―─ｰ—'})


def norm_code(s):
    """品番/商品コードを比較しやすいように正規化する。
    OCRはハイフンを全角ダッシュ・長音記号・漢字の「一」などに
    誤認識しやすく、まれに連続したハイフンとして重複読みすることも
    あるため、ダッシュ類は半角ハイフン1つに統一する。
    """
    if s is None:
        return ''
    s = str(s).translate(DASH_CHARS)
    s = re.sub(r'-{2,}', '-', s)
    s = re.sub(r'\s+', '', s)
    return s.strip().upper()


def code_similarity(a, b):
    """0.0〜1.0 のコード類似度。OCR誤読による軽微なズレを吸収する。"""
    if not a or not b:
        return 0.0
    return difflib.SequenceMatcher(None, a, b).ratio()


# ---------------------------------------------------------------------------
# PDF OCR 抽出
# ---------------------------------------------------------------------------

@dataclass
class PdfRow:
    page: int
    rowno: int
    code: str
    name: str
    qty: int
    qty_raw: str
    suspicious: bool
    wh: str = ''


HEADER_LABELS = {
    'rowno': '行番号',
    'code': '商品コード',
    'name': '商品名',
    'qty': '移動数',
    'wh': '出庫倉庫コード',
    'remark': '移動明細摘要',
}


def _cluster_rows(words, y_tol):
    words = sorted(words, key=lambda w: (w['y'], w['x']))
    rows = []
    for w in words:
        placed = False
        for row in rows:
            if w['y'] < row['ymax'] and w['y'] + w['h'] > row['ymin']:
                row['words'].append(w)
                row['ymin'] = min(row['ymin'], w['y'])
                row['ymax'] = max(row['ymax'], w['y'] + w['h'])
                placed = True
                break
        if not placed:
            rows.append({'words': [w], 'ymin': w['y'], 'ymax': w['y'] + w['h']})
    rows.sort(key=lambda r: r['ymin'])
    for row in rows:
        row['words'].sort(key=lambda w: w['x'])
    return rows


def _find_header_boundaries(rows):
    for row in rows:
        full = ''.join(w['text'] for w in row['words'])
        if '行番号' in full and '商品コード' in full:
            joined_chars = []
            for w in row['words']:
                for ch in w['text']:
                    joined_chars.append((ch, w['x']))
            full2 = ''.join(c for c, _ in joined_chars)
            starts = {}
            for key, label in HEADER_LABELS.items():
                idx = full2.find(label)
                if idx >= 0:
                    starts[key] = joined_chars[idx][1]
            if len(starts) >= 4:
                return starts
    return None


def _bucket_for(x, bounds):
    order = ['rowno', 'code', 'name', 'qty', 'wh', 'remark']
    present = sorted(((k, bounds[k]) for k in order if k in bounds), key=lambda kv: kv[1])
    best = present[0][0]
    for k, start in present:
        if x >= start - 20:
            best = k
        else:
            break
    return best


async def _ocr_page(png_path, lang):
    img = Image.open(png_path)
    result = await winocr.recognize_pil(img, lang=lang)
    words = []
    for line in result.lines:
        for w in line.words:
            r = w.bounding_rect
            words.append({'text': w.text, 'x': r.x, 'y': r.y, 'w': r.width, 'h': r.height})
    return words


# 数量セルの再読み取り(検算)倍率。全ページ一括OCRだと「00」のような
# 数字の組み合わせがまれに1文字の漢字風グリフに誤認識されることがあるため、
# 数量セルだけを切り出して高解像度で再OCRし、読み取り精度を底上げする。
QTY_VERIFY_SCALE = 12


async def _reocr_qty_cell(page, base_scale, x0, y0, x1, y1, lang):
    """数量セルをbase_scaleのピクセル座標で受け取り、高解像度で再OCRして
    数字だけの文字列を返す(読み取れない場合は None)。"""
    rect = fitz.Rect(x0 / base_scale, y0 / base_scale, x1 / base_scale, y1 / base_scale)
    pix = page.get_pixmap(matrix=fitz.Matrix(QTY_VERIFY_SCALE, QTY_VERIFY_SCALE), clip=rect)
    img = Image.frombytes('RGB', (pix.width, pix.height), pix.samples)
    result = await winocr.recognize_pil(img, lang=lang)
    text = result.text
    if re.sub(r'[0-9,.\s]', '', text):
        return None  # 数字・カンマ以外の文字が残る = 読み取り不確実
    digits = re.sub(r'[^0-9]', '', text)
    return digits or None


async def extract_pdf_rows_async(pdf_path, scale=3, lang='ja', progress_cb=None):
    """PDF(移動伝票)をOCRして明細行のリストを返す。"""
    if winocr is None:
        raise RuntimeError(
            "winocr がインストールされていません。\n"
            "コマンドプロンプトで `pip install winocr` を実行してください。"
        )
    doc = fitz.open(pdf_path)
    all_rows = []
    last_bounds = None
    n_pages = len(doc)
    for pno in range(n_pages):
        if progress_cb:
            progress_cb(pno, n_pages, f"PDF {pno + 1}/{n_pages} ページ OCR中...")
        page = doc[pno]
        pix = page.get_pixmap(matrix=fitz.Matrix(scale, scale))
        tmp_png = Path(pdf_path).with_name(f"_ocr_tmp_page_{pno}.png")
        pix.save(str(tmp_png))
        try:
            words = await _ocr_page(str(tmp_png), lang)
        finally:
            try:
                tmp_png.unlink()
            except OSError:
                pass
        y_tol = scale * 5
        rows = _cluster_rows(words, y_tol)
        bounds = _find_header_boundaries(rows)
        if bounds is not None:
            last_bounds = bounds
        elif last_bounds is not None:
            bounds = last_bounds
        else:
            continue  # このページはテーブル形式を認識できなかった

        for row in rows:
            cols = {'rowno': [], 'code': [], 'name': [], 'qty': [], 'wh': [], 'remark': []}
            for w in row['words']:
                cols[_bucket_for(w['x'], bounds)].append(w['text'])
            rowno_txt = ''.join(cols['rowno']).replace(' ', '')
            if not re.fullmatch(r'\d{1,4}', rowno_txt):
                continue
            code = norm_code(''.join(cols['code']))
            name = ''.join(cols['name'])
            qty_raw = ''.join(cols['qty'])
            qty_digits = re.sub(r'[^0-9]', '', qty_raw)
            suspicious = bool(re.sub(r'[0-9,\s]', '', qty_raw))
            qty = int(qty_digits) if qty_digits else None
            wh = ''.join(cols['wh']).replace(' ', '')
            if qty is None:
                suspicious = True
                qty = 0

            # 数量セルを高解像度で再OCRして検算する(誤読しやすい列のため)
            if 'qty' in bounds and 'wh' in bounds:
                x0 = bounds['qty'] - 15
                x1 = bounds['wh'] - 10
                y0 = row['ymin'] - scale * 3
                y1 = row['ymax'] + scale * 3
                try:
                    hi_digits = await _reocr_qty_cell(page, scale, x0, y0, x1, y1, lang)
                except Exception:
                    hi_digits = None
                if hi_digits:
                    hi_val = int(hi_digits)
                    if suspicious:
                        qty = hi_val
                        suspicious = False
                    elif hi_val != qty:
                        suspicious = True
                        qty_raw = f'{qty_raw} (再読取結果:{hi_digits})'

            all_rows.append(PdfRow(
                page=pno + 1, rowno=int(rowno_txt), code=code, name=name,
                qty=qty, qty_raw=qty_raw, suspicious=suspicious, wh=wh,
            ))
    if progress_cb:
        progress_cb(n_pages, n_pages, "PDF OCR完了")
    return all_rows


def extract_pdf_rows(pdf_path, scale=3, lang='ja', progress_cb=None):
    return asyncio.run(extract_pdf_rows_async(pdf_path, scale=scale, lang=lang, progress_cb=progress_cb))


# ---------------------------------------------------------------------------
# Excel 読み取り
# ---------------------------------------------------------------------------

@dataclass
class ExcelRow:
    row_idx: int
    section: str
    code: str
    name: str
    qty: int
    note: str = ''


def list_sheet_names(xlsx_path):
    wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
    try:
        return wb.sheetnames
    finally:
        wb.close()


def _is_blank_row(values):
    return all(v is None or (isinstance(v, str) and v.strip() == '') for v in values)


def extract_excel_rows(xlsx_path, sheet_name):
    """指定シートから 品番/品名/支給数 のある明細行だけを抽出する。
    「簡易品名」等、品番を持たないセクション(補材/返却品など)は対象外。
    """
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    ws = wb[sheet_name]
    header_map = None  # {col_index(0-based): label}
    rows_out = []
    current_section = ''
    for r_idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
        values = list(row)
        if _is_blank_row(values):
            header_map = None
            continue
        first = str(values[0]).strip() if values[0] is not None else ''
        # ヘッダー行の検出 (例: 品番/品名/支給数/パレット番号/複数パレット/備考)
        norm_vals = [str(v).strip() if v is not None else '' for v in values]
        if first in ('品番', '簡易品名') and ('品名' in norm_vals or '簡易品名' in norm_vals):
            header_map = {i: v for i, v in enumerate(norm_vals) if v}
            continue
        if header_map is None:
            # ヘッダーが見つかっていない行 = セクション見出し or その他の行
            if first and all(v is None for v in values[1:]):
                current_section = first
            continue
        # データ行
        code_col = None
        for i, label in header_map.items():
            if label == '品番':
                code_col = i
                break
        if code_col is None:
            # 品番列を持たないセクション(簡易品名など)は照合対象外
            continue
        code_val = values[code_col] if code_col < len(values) else None
        if code_val is None or str(code_val).strip() == '':
            continue
        name_col = None
        qty_col = None
        for i, label in header_map.items():
            if label == '品名':
                name_col = i
            if label in ('支給数', '数量'):
                qty_col = i
        name_val = values[name_col] if (name_col is not None and name_col < len(values)) else ''
        qty_val = values[qty_col] if (qty_col is not None and qty_col < len(values)) else None
        try:
            qty_int = int(qty_val)
        except (TypeError, ValueError):
            qty_int = None
        note_parts = []
        for i, label in header_map.items():
            if label in ('備考',) and i < len(values) and values[i]:
                note_parts.append(str(values[i]))
        rows_out.append(ExcelRow(
            row_idx=r_idx,
            section=current_section,
            code=norm_code(code_val),
            name=str(name_val) if name_val is not None else '',
            qty=qty_int if qty_int is not None else 0,
            note='; '.join(note_parts),
        ))
    wb.close()
    return rows_out


# ---------------------------------------------------------------------------
# 照合ロジック
# ---------------------------------------------------------------------------

@dataclass
class MatchResult:
    status: str            # '一致' / '数量相違' / '要確認(OCR)' / 'PDFに見つかりません' / 'Excelに見つかりません'
    excel_row: 'ExcelRow' = None
    pdf_rows: list = field(default_factory=list)
    excel_code: str = ''
    pdf_code: str = ''
    excel_name: str = ''
    pdf_name: str = ''
    excel_qty: int = None
    pdf_qty: int = None
    note: str = ''

    @property
    def diff(self):
        if self.excel_qty is None or self.pdf_qty is None:
            return None
        return self.pdf_qty - self.excel_qty

    @property
    def pdf_rowno(self):
        """PDF(移動伝票)側の行番号。該当行が無ければ空文字、複数行が
        1つのExcel行に対応する場合(例: 同一品番が2行に分かれて記載)は
        カンマ区切りで併記する。原本PDFで該当箇所をすぐ探せるようにするため。
        """
        if not self.pdf_rows:
            return ''
        return ', '.join(str(p.rowno) for p in sorted(self.pdf_rows, key=lambda p: p.rowno))


FUZZY_THRESHOLD = 0.72


def compare(excel_rows, pdf_rows):
    """Excel行を基準に、PDF行(OCR抽出結果)と突き合わせる。"""
    pdf_pool = list(pdf_rows)  # 消費していく候補プール
    results = []

    # 1) 完全一致(正規化コード一致)をすべて集約
    remaining_excel = []
    for erow in excel_rows:
        exact = [p for p in pdf_pool if p.code == erow.code]
        if exact:
            for p in exact:
                pdf_pool.remove(p)
            total_qty = sum(p.qty for p in exact)
            any_suspicious = any(p.suspicious for p in exact)
            names = ' / '.join(sorted(set(p.name for p in exact)))
            if any_suspicious:
                status = '要確認(OCR)'
                note = 'PDF側の数量読み取りに自信が持てない行があります。目視で確認してください。'
            elif total_qty == erow.qty:
                status = '一致'
                note = ''
            else:
                status = '数量相違'
                note = ''
            results.append(MatchResult(
                status=status, excel_row=erow, pdf_rows=exact,
                excel_code=erow.code, pdf_code=exact[0].code,
                excel_name=erow.name, pdf_name=names,
                excel_qty=erow.qty, pdf_qty=total_qty, note=note,
            ))
        else:
            remaining_excel.append(erow)

    # 2) 残りはあいまい一致(OCR誤読を吸収するためのコード類似度)
    still_unmatched_excel = []
    for erow in remaining_excel:
        best = None
        best_score = 0.0
        for p in pdf_pool:
            score = code_similarity(erow.code, p.code)
            if score > best_score:
                best_score = score
                best = p
        if best is not None and best_score >= FUZZY_THRESHOLD:
            pdf_pool.remove(best)
            status = '要確認(OCR)'
            note = f'品番のOCR読み取りが完全一致しません(類似度{best_score:.0%})。目視で確認してください。'
            if best.suspicious:
                note += ' 数量読み取りも不明瞭です。'
            results.append(MatchResult(
                status=status, excel_row=erow, pdf_rows=[best],
                excel_code=erow.code, pdf_code=best.code,
                excel_name=erow.name, pdf_name=best.name,
                excel_qty=erow.qty, pdf_qty=best.qty, note=note,
            ))
        else:
            still_unmatched_excel.append(erow)

    # 3) Excel側に対応するPDF行が見つからなかったもの
    for erow in still_unmatched_excel:
        results.append(MatchResult(
            status='PDFに見つかりません', excel_row=erow,
            excel_code=erow.code, excel_name=erow.name, excel_qty=erow.qty,
            pdf_code='', pdf_name='', pdf_qty=None,
            note='PDF(支給部材)側にこの品番が見つかりませんでした。',
        ))

    # 4) PDF側で余ったもの(Excelに対応行がない)
    for p in pdf_pool:
        results.append(MatchResult(
            status='Excelに見つかりません', pdf_rows=[p],
            excel_code='', excel_name='', excel_qty=None,
            pdf_code=p.code, pdf_name=p.name, pdf_qty=p.qty,
            note=f'PDF {p.page}ページ {p.rowno}行目。Excel(移動明細)側に対応する品番が見つかりませんでした。'
                 + (' 数量読み取り不明瞭。' if p.suspicious else ''),
        ))

    return results


STATUS_ORDER = {'数量相違': 0, '要確認(OCR)': 1, 'PDFに見つかりません': 2,
                'Excelに見つかりません': 3, '一致': 4}


def _sort_key(res):
    """確認が必要な状態を上に、それ以外(グループ内)は行番号順に並べるキー。"""
    group = STATUS_ORDER.get(res.status, 9)
    if res.pdf_rows:
        rowno = min(p.rowno for p in res.pdf_rows)
    elif res.excel_row is not None:
        rowno = res.excel_row.row_idx
    else:
        rowno = 0
    return (group, rowno)


def sort_results(results):
    return sorted(results, key=_sort_key)


# ---------------------------------------------------------------------------
# Excel レポート出力
# ---------------------------------------------------------------------------

import unicodedata

NOTE_COL_WIDTH = 60  # 備考列の幅(Excel列幅単位)


def _display_width(s):
    """全角文字を2、半角文字を1として数えた表示幅。"""
    if not s:
        return 0
    return sum(2 if unicodedata.east_asian_width(ch) in ('F', 'W', 'A') else 1 for ch in str(s))


def _estimate_row_height(text, col_width, base_height=15, chars_per_line_margin=2):
    """折り返し後の行数を概算し、備考(全文表示)が見切れないよう行の高さを見積もる。"""
    if not text:
        return base_height
    usable = max(1, col_width - chars_per_line_margin)
    total_lines = 0
    for line in str(text).split('\n'):
        width = _display_width(line)
        total_lines += max(1, -(-width // usable))  # 切り上げ除算
    return max(base_height, base_height * total_lines)


FILL_OK = PatternFill('solid', fgColor='C6EFCE')
FILL_NG = PatternFill('solid', fgColor='FFC7CE')
FILL_WARN = PatternFill('solid', fgColor='FFEB9C')
FILL_MISSING = PatternFill('solid', fgColor='D9D9D9')

STATUS_FILL = {
    '一致': FILL_OK,
    '数量相違': FILL_NG,
    '要確認(OCR)': FILL_WARN,
    'PDFに見つかりません': FILL_MISSING,
    'Excelに見つかりません': FILL_MISSING,
}


def write_report(results, out_path, pdf_name, sheet_name):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = '照合結果'

    ws['A1'] = '支給部材PDF ⇔ 移動明細Excel 照合結果'
    ws['A1'].font = Font(bold=True, size=14)
    ws['A2'] = f'PDF: {pdf_name}'
    ws['A3'] = f'Excelシート: {sheet_name}'

    headers = ['PDF行番号', '状態', 'Excel品番', 'PDF品番(OCR)', 'Excel品名', 'PDF品名(OCR参考)',
               'Excel支給数', 'PDF移動数(OCR)', '差異', '備考']
    note_col = len(headers)
    header_row = 5
    for c, h in enumerate(headers, start=1):
        cell = ws.cell(row=header_row, column=c, value=h)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal='center')

    results_sorted = sort_results(results)

    r = header_row + 1
    counts = {}
    for res in results_sorted:
        counts[res.status] = counts.get(res.status, 0) + 1
        diff = res.diff
        values = [
            res.pdf_rowno, res.status, res.excel_code, res.pdf_code, res.excel_name, res.pdf_name,
            res.excel_qty, res.pdf_qty, diff if diff is not None else '', res.note,
        ]
        fill = STATUS_FILL.get(res.status)
        for c, v in enumerate(values, start=1):
            cell = ws.cell(row=r, column=c, value=v)
            if fill:
                cell.fill = fill
            if c == note_col:
                cell.alignment = Alignment(wrap_text=True, vertical='top')
        ws.row_dimensions[r].height = _estimate_row_height(res.note, NOTE_COL_WIDTH)
        r += 1

    widths = [11, 16, 16, 16, 26, 26, 12, 14, 8, NOTE_COL_WIDTH]
    for c, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(c)].width = w

    # サマリーシート
    ws2 = wb.create_sheet('サマリー')
    ws2['A1'] = '照合結果サマリー'
    ws2['A1'].font = Font(bold=True, size=13)
    ws2['A2'] = f'PDF: {pdf_name}'
    ws2['A3'] = f'Excelシート: {sheet_name}'
    row = 5
    ws2.cell(row=row, column=1, value='状態').font = Font(bold=True)
    ws2.cell(row=row, column=2, value='件数').font = Font(bold=True)
    row += 1
    for status in ['一致', '数量相違', '要確認(OCR)', 'PDFに見つかりません', 'Excelに見つかりません']:
        ws2.cell(row=row, column=1, value=status)
        ws2.cell(row=row, column=2, value=counts.get(status, 0))
        fill = STATUS_FILL.get(status)
        if fill:
            ws2.cell(row=row, column=1).fill = fill
            ws2.cell(row=row, column=2).fill = fill
        row += 1
    ws2.column_dimensions['A'].width = 22
    ws2.column_dimensions['B'].width = 10

    wb.save(out_path)
    return counts


# ---------------------------------------------------------------------------
# GUI
# ---------------------------------------------------------------------------

STATUS_TAG_COLORS = {
    '一致': '#C6EFCE',
    '数量相違': '#FFC7CE',
    '要確認(OCR)': '#FFEB9C',
    'PDFに見つかりません': '#D9D9D9',
    'Excelに見つかりません': '#D9D9D9',
}


def guess_sheet_from_filename(pdf_path, sheet_names):
    m = re.search(r'(\d{6})', Path(pdf_path).stem)
    if not m:
        return None
    digits = m.group(1)
    yy, mm, dd = digits[0:2], digits[2:4], digits[4:6]
    mm_i, dd_i = str(int(mm)), str(int(dd))
    candidates = [f'{yy}.{mm_i}.{dd_i}', f'{yy}.{mm}.{dd}']
    for cand in candidates:
        for name in sheet_names:
            if name.startswith(cand):
                return name
    return None


class App:
    def __init__(self, root):
        self.root = root
        root.title('支給部材PDF ⇔ 移動明細Excel 照合ツール   Ver %s' % get_version())
        root.geometry('1180x680')

        self.pdf_path = StringVar()
        self.xlsx_path = StringVar()
        self.sheet_name = StringVar()
        self.status_text = StringVar(value='PDFとExcelを選択してください。')
        self.results = []
        self.last_report_path = None
        self.config = load_config()

        self._build_widgets()

    # -- widgets -------------------------------------------------------
    def _build_widgets(self):
        pad = {'padx': 8, 'pady': 6}

        frm_top = ttk.Frame(self.root)
        frm_top.pack(side=TOP, fill=X, **pad)

        ttk.Label(frm_top, text='支給部材PDF:').grid(row=0, column=0, sticky=W)
        ttk.Entry(frm_top, textvariable=self.pdf_path, width=90).grid(row=0, column=1, sticky=W + E, padx=4)
        ttk.Button(frm_top, text='参照...', command=self.pick_pdf).grid(row=0, column=2)

        ttk.Label(frm_top, text='移動明細Excel:').grid(row=1, column=0, sticky=W, pady=(6, 0))
        ttk.Entry(frm_top, textvariable=self.xlsx_path, width=90).grid(row=1, column=1, sticky=W + E, padx=4, pady=(6, 0))
        ttk.Button(frm_top, text='参照...', command=self.pick_xlsx).grid(row=1, column=2, pady=(6, 0))

        ttk.Label(frm_top, text='対象シート:').grid(row=2, column=0, sticky=W, pady=(6, 0))
        self.sheet_combo = ttk.Combobox(frm_top, textvariable=self.sheet_name, width=40, state='readonly')
        self.sheet_combo.grid(row=2, column=1, sticky=W, padx=4, pady=(6, 0))

        frm_top.columnconfigure(1, weight=1)

        frm_btn = ttk.Frame(self.root)
        frm_btn.pack(side=TOP, fill=X, **pad)
        self.run_btn = ttk.Button(frm_btn, text='照合実行', command=self.run_check)
        self.run_btn.pack(side=LEFT)
        self.export_btn = ttk.Button(frm_btn, text='結果をExcelに出力...', command=self.export_report, state='disabled')
        self.export_btn.pack(side=LEFT, padx=8)
        ttk.Button(frm_btn, text='使い方', command=self._open_manual).pack(side=RIGHT, padx=(0, 6))
        ttk.Label(frm_btn, textvariable=self.status_text).pack(side=LEFT, padx=16)

        self.progress = ttk.Progressbar(self.root, orient=HORIZONTAL, mode='determinate')
        self.progress.pack(side=TOP, fill=X, padx=8)

        # 結果テーブル
        columns = ('pdf_rowno', 'status', 'excel_code', 'pdf_code', 'excel_name', 'pdf_name', 'excel_qty', 'pdf_qty', 'diff', 'note')
        headers = ['行№', '状態', 'Excel品番', 'PDF品番(OCR)', 'Excel品名', 'PDF品名(OCR参考)', 'Excel支給数', 'PDF移動数(OCR)', '差異', '備考']
        widths = [60, 110, 120, 120, 220, 220, 90, 100, 60, 320]
        anchors = {'pdf_rowno': E}

        frm_table = ttk.Frame(self.root)
        frm_table.pack(side=TOP, fill=BOTH, expand=True, padx=8, pady=(4, 8))

        self.tree = ttk.Treeview(frm_table, columns=columns, show='headings')
        for col, h, w in zip(columns, headers, widths):
            self.tree.heading(col, text=h)
            self.tree.column(col, width=w, anchor=anchors.get(col, W))

        vsb = ttk.Scrollbar(frm_table, orient=VERTICAL, command=self.tree.yview)
        hsb = ttk.Scrollbar(frm_table, orient=HORIZONTAL, command=self.tree.xview)
        self.tree.configure(yscrollcommand=vsb.set, xscrollcommand=hsb.set)

        self.tree.grid(row=0, column=0, sticky='nsew')
        vsb.grid(row=0, column=1, sticky='ns')
        hsb.grid(row=1, column=0, sticky='ew')
        frm_table.rowconfigure(0, weight=1)
        frm_table.columnconfigure(0, weight=1)

        for status, color in STATUS_TAG_COLORS.items():
            self.tree.tag_configure(status, background=color)

    # -- file pickers ----------------------------------------------------
    def _remember_dir(self, key, file_path):
        self.config[key] = str(Path(file_path).parent)
        save_config(self.config)

    def pick_pdf(self):
        path = filedialog.askopenfilename(
            title='支給部材PDFを選択', filetypes=[('PDF', '*.pdf')],
            initialdir=self.config.get('pdf_dir', ''))
        if path:
            self.pdf_path.set(path)
            self._remember_dir('pdf_dir', path)
            if self.xlsx_path.get():
                self._refresh_sheet_guess()

    def pick_xlsx(self):
        path = filedialog.askopenfilename(
            title='移動明細Excelを選択', filetypes=[('Excel', '*.xlsx *.xlsm')],
            initialdir=self.config.get('xlsx_dir', ''))
        if path:
            self.xlsx_path.set(path)
            self._remember_dir('xlsx_dir', path)
            try:
                names = list_sheet_names(path)
            except Exception as e:
                messagebox.showerror('エラー', f'Excelの読み込みに失敗しました:\n{e}')
                return
            self.sheet_combo['values'] = names
            if names:
                self.sheet_name.set(names[-1])
            if self.pdf_path.get():
                self._refresh_sheet_guess()

    def _refresh_sheet_guess(self):
        names = list(self.sheet_combo['values'])
        if not names:
            return
        guess = guess_sheet_from_filename(self.pdf_path.get(), names)
        if guess:
            self.sheet_name.set(guess)

    # -- run ---------------------------------------------------------
    def run_check(self):
        pdf = self.pdf_path.get().strip()
        xlsx = self.xlsx_path.get().strip()
        sheet = self.sheet_name.get().strip()
        if not pdf or not Path(pdf).exists():
            messagebox.showwarning('確認', 'PDFファイルを選択してください。')
            return
        if not xlsx or not Path(xlsx).exists():
            messagebox.showwarning('確認', 'Excelファイルを選択してください。')
            return
        if not sheet:
            messagebox.showwarning('確認', '対象シートを選択してください。')
            return

        self.run_btn.config(state='disabled')
        self.export_btn.config(state='disabled')
        self.tree.delete(*self.tree.get_children())
        self.progress['value'] = 0
        self.status_text.set('照合処理を開始します...')

        thread = threading.Thread(target=self._run_check_worker, args=(pdf, xlsx, sheet), daemon=True)
        thread.start()

    def _progress_cb(self, done, total, message):
        def upd():
            self.progress['maximum'] = max(total, 1)
            self.progress['value'] = done
            self.status_text.set(message)
        self.root.after(0, upd)

    def _run_check_worker(self, pdf, xlsx, sheet):
        try:
            pdf_rows = extract_pdf_rows(pdf, progress_cb=self._progress_cb)
            self._progress_cb(1, 1, 'Excelを読み込み中...')
            excel_rows = extract_excel_rows(xlsx, sheet)
            self._progress_cb(1, 1, '照合中...')
            results = compare(excel_rows, pdf_rows)
        except Exception as e:
            traceback.print_exc()
            self.root.after(0, lambda: self._on_error(e))
            return
        self.root.after(0, lambda: self._on_done(results))

    def _on_error(self, e):
        self.run_btn.config(state='normal')
        self.status_text.set('エラーが発生しました。')
        messagebox.showerror('エラー', str(e))

    def _on_done(self, results):
        self.results = results
        results_sorted = sort_results(results)
        for res in results_sorted:
            diff = res.diff
            self.tree.insert('', END, values=(
                res.pdf_rowno, res.status, res.excel_code, res.pdf_code, res.excel_name, res.pdf_name,
                res.excel_qty if res.excel_qty is not None else '',
                res.pdf_qty if res.pdf_qty is not None else '',
                diff if diff is not None else '',
                res.note,
            ), tags=(res.status,))

        counts = {}
        for res in results:
            counts[res.status] = counts.get(res.status, 0) + 1
        summary = ' / '.join(f'{k}:{v}' for k, v in counts.items())
        self.status_text.set(f'完了。 {summary}')
        self.run_btn.config(state='normal')
        self.export_btn.config(state='normal')

    def export_report(self):
        if not self.results:
            return
        default_name = f"照合結果_{Path(self.pdf_path.get()).stem}.xlsx"
        out_path = filedialog.asksaveasfilename(
            title='照合結果の保存先', defaultextension='.xlsx',
            initialfile=default_name, filetypes=[('Excel', '*.xlsx')],
            initialdir=self.config.get('export_dir', ''))
        if not out_path:
            return
        try:
            write_report(self.results, out_path, Path(self.pdf_path.get()).name, self.sheet_name.get())
        except Exception as e:
            messagebox.showerror('エラー', f'出力に失敗しました:\n{e}')
            return
        self._remember_dir('export_dir', out_path)
        self.last_report_path = out_path
        if messagebox.askyesno('完了', f'{out_path}\nに出力しました。開きますか?'):
            import os
            os.startfile(out_path)

    # -- 使い方 ---------------------------------------------------------
    def _open_manual(self):
        win = Toplevel(self.root)
        win.title('使い方   Ver %s' % get_version())
        win.geometry('700x600')
        win.transient(self.root)

        frame = ttk.Frame(win)
        frame.pack(fill=BOTH, expand=True)
        text = tk.Text(frame, wrap=WORD, padx=10, pady=10)
        text.pack(side=LEFT, fill=BOTH, expand=True)
        sb = ttk.Scrollbar(frame, orient=VERTICAL, command=text.yview)
        sb.pack(side=LEFT, fill=Y)
        text.configure(yscrollcommand=sb.set)

        try:
            with open(_resource_path('使い方.md'), encoding='utf-8') as f:
                content = f.read()
        except OSError as e:
            content = f'使い方.md の読み込みに失敗しました。\n\n{e}'
        text.insert('1.0', content)
        text.configure(state='disabled')

        ttk.Button(win, text='閉じる', command=win.destroy).pack(pady=(0, 10))


def main():
    root = Tk()
    App(root)
    root.mainloop()


if __name__ == '__main__':
    main()
