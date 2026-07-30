# -*- coding: utf-8 -*-
"""
製品構成 逆検索ツール（スタンドアロンGUI版）
  Excel/VBA版 製品構成逆検索ツール Ver1.04 を Python(Tkinter) に移植したもの。

機能:
  (起動時に自動)        : TPiCS(SQL Server) から XITEM/XPRTS/XHEAD/XSECT を自動読込
  (1) 機種コード逆展開  : 部品 → それを使う最上位の機種コード・使用数・客先品目
  (2) 直上親コード検索  : 部品 → 直上親(TD-仮親はスキップ)・名称・代表機種・メーカー
  (3) 別ファイル保存    : 検索結果を CSV / Excel で保存

DBへは SELECT のみ（書込みは行わない）。tdread は読取り専用ユーザー。
"""

import os
import sys
import csv
import threading
import configparser
import datetime
import tkinter as tk
from tkinter import ttk, messagebox, filedialog

import core
import dbsource

if getattr(sys, "frozen", False):
    # PyInstaller等でexe化した場合は exe と同じフォルダの config.ini を読む
    BASE = os.path.dirname(sys.executable)
else:
    BASE = os.path.dirname(os.path.abspath(__file__))

cfg = configparser.ConfigParser()
cfg.read(os.path.join(BASE, "config.ini"), encoding="utf-8")

DEMO = ("--demo" in sys.argv) or (
    cfg.get("app", "demo", fallback="false").lower() == "true")


def db_cfg():
    if not cfg.has_section("db"):
        return {}
    return dict(cfg.items("db"))


class App(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title("製品構成 逆検索ツール" + ("（デモモード）" if DEMO else ""))
        self.geometry("1180x720")
        self.minsize(900, 560)

        self.index = None          # core.BomIndex
        self.sync_time = None
        self.current_rows = []     # 表示中の結果(dict)
        self.current_mode = None   # "kishu" / "chokujou"

        self._build_ui()

        # 起動時：マスタを自動読込（デモ=内蔵データ / 本番=さくらクラウドのSQL Server）
        self.after(150, self._auto_load_master)

    # ----------------------------------------------------------------- UI
    def _build_ui(self):
        # 上部ツールバー
        top = ttk.Frame(self, padding=(10, 8))
        top.pack(side=tk.TOP, fill=tk.X)

        self.btn_kishu = ttk.Button(top, text="① 機種コード逆展開検索",
                                    command=lambda: self.on_search("kishu"))
        self.btn_kishu.pack(side=tk.LEFT, padx=(0, 6))
        self.btn_choku = ttk.Button(top, text="② 直上親コード検索",
                                    command=lambda: self.on_search("chokujou"))
        self.btn_choku.pack(side=tk.LEFT)

        ttk.Separator(top, orient=tk.VERTICAL).pack(
            side=tk.LEFT, fill=tk.Y, padx=10)

        self.btn_save = ttk.Button(top, text="③ 別ファイル保存",
                                   command=self.on_save)
        self.btn_save.pack(side=tk.LEFT)
        # マスタは起動時に自動読込（再読込ボタンは無し。最新化はアプリ再起動）。

        # 本体（左:入力 / 右:結果）
        main = ttk.Panedwindow(self, orient=tk.HORIZONTAL)
        main.pack(side=tk.TOP, fill=tk.BOTH, expand=True, padx=10, pady=(0, 6))

        left = ttk.Labelframe(main, text="検索する部品コード（1行に1コード）",
                              padding=6)
        self.txt = tk.Text(left, width=24, undo=True, font=("Consolas", 11))
        scl = ttk.Scrollbar(left, command=self.txt.yview)
        self.txt.configure(yscrollcommand=scl.set)
        self.txt.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        scl.pack(side=tk.LEFT, fill=tk.Y)
        btnclr = ttk.Button(left, text="入力クリア", command=self.on_clear)
        btnclr.pack(side=tk.BOTTOM, fill=tk.X, pady=(6, 0))
        main.add(left, weight=0)

        right = ttk.Labelframe(main, text="検索結果", padding=6)
        self.tree = ttk.Treeview(right, show="headings", selectmode="extended")
        ysb = ttk.Scrollbar(right, orient=tk.VERTICAL, command=self.tree.yview)
        xsb = ttk.Scrollbar(right, orient=tk.HORIZONTAL, command=self.tree.xview)
        self.tree.configure(yscrollcommand=ysb.set, xscrollcommand=xsb.set)
        self.tree.grid(row=0, column=0, sticky="nsew")
        ysb.grid(row=0, column=1, sticky="ns")
        xsb.grid(row=1, column=0, sticky="ew")
        right.rowconfigure(0, weight=1)
        right.columnconfigure(0, weight=1)
        self.tree.tag_configure("odd", background="#f4f8fc")
        main.add(right, weight=1)

        # ステータスバー
        self.status = tk.StringVar()
        bar = ttk.Frame(self, relief=tk.SUNKEN, padding=(8, 3))
        bar.pack(side=tk.BOTTOM, fill=tk.X)
        ttk.Label(bar, textvariable=self.status, anchor="w").pack(
            side=tk.LEFT, fill=tk.X, expand=True)
        self.badge = ttk.Label(bar, text="", anchor="e")
        self.badge.pack(side=tk.RIGHT)

        self._set_status("起動中...")
        self._update_badge()
        self._set_search_enabled(False)

    def _update_badge(self):
        if DEMO:
            self.badge.config(text="● デモモード(DB未接続)", foreground="#c0392b")
        elif self.index is not None:
            self.badge.config(text="● 本番(マスタ取得済)", foreground="#1e8449")
        else:
            self.badge.config(text="○ 本番(マスタ未取得)", foreground="#7f8c8d")

    def _set_status(self, msg):
        self.status.set(msg)

    def _set_search_enabled(self, on):
        state = tk.NORMAL if on else tk.DISABLED
        self.btn_kishu.config(state=state)
        self.btn_choku.config(state=state)
        self.btn_save.config(state=state)

    # --------------------------------------------------- 非同期ヘルパ
    def _run_async(self, fn, on_done=None, on_error=None):
        """fn を別スレッドで実行し、完了時に on_done/on_error をUIスレッドで呼ぶ。"""
        self._set_busy(True)

        def worker():
            try:
                result = fn()
            except Exception as e:
                self.after(0, lambda: self._async_fail(e, on_error))
            else:
                self.after(0, lambda: self._async_ok(result, on_done))
        threading.Thread(target=worker, daemon=True).start()

    def _async_ok(self, result, on_done):
        self._set_busy(False)
        if on_done:
            on_done(result)

    def _async_fail(self, err, on_error):
        self._set_busy(False)
        if on_error:
            on_error(err)
        else:
            self._set_status("エラー: %s" % err)
            messagebox.showerror("エラー", str(err))

    def _set_busy(self, busy):
        self.config(cursor="watch" if busy else "")
        if busy:
            self._set_search_enabled(False)
        else:
            self._set_search_enabled(self.index is not None)

    # --------------------------------------------------- マスタ読込
    def _load_demo(self):
        return dbsource.load_index_demo(progress=self._progress)

    def _load_db(self):
        return dbsource.load_index_from_db(db_cfg(), progress=self._progress)

    def _progress(self, msg):
        # 別スレッドからの進捗。UIスレッドへ委譲。
        self.after(0, lambda: self._set_status(msg))

    def _auto_load_master(self):
        """起動時の自動読込。"""
        self._load_master()

    def _load_master(self):
        if DEMO:
            self._set_status("デモデータを読み込んでいます...")
            self._run_async(self._load_demo, on_done=self._after_master_loaded,
                            on_error=self._after_master_error)
            return
        self._set_status("さくらクラウドのSQL Serverへ接続してマスタを読込中...（最大2分）")
        self._run_async(self._load_db, on_done=self._after_master_loaded,
                        on_error=self._after_master_error)

    def _after_master_loaded(self, index):
        self.index = index
        self.sync_time = datetime.datetime.now()
        n_item = len(index.info_bunr)
        n_prts = sum(len(v) for v in index.dict_fwd.values())
        self._update_badge()
        self._set_search_enabled(True)
        self._set_status("マスタ読込完了  %s  （品目 %d件 / 構成 %d行）"
                         % (self.sync_time.strftime("%Y/%m/%d %H:%M:%S"),
                            n_item, n_prts))

    def _after_master_error(self, err):
        self.index = None
        self._update_badge()
        self._set_status("マスタ取得失敗: %s" % err)
        messagebox.showerror(
            "接続エラー",
            "接続に失敗またはタイムアウトしました。\n"
            "通信環境（VPNなど）と config.ini の接続情報を確認してください。\n\n詳細: %s" % err)

    # --------------------------------------------------- 検索
    def _input_parts(self):
        text = self.txt.get("1.0", tk.END)
        seen = []
        for line in text.splitlines():
            c = line.strip()
            if c:
                seen.append(c)
        return seen

    def on_search(self, mode):
        if self.index is None:
            messagebox.showwarning(
                "確認", "マスタを読込中、または読込に失敗しています。\n"
                "数秒待つか、アプリを再起動してください。")
            return
        parts = self._input_parts()
        if not parts:
            messagebox.showwarning(
                "確認", "検索する部品コードを入力してください（1行に1コード）。")
            return
        self.current_mode = mode
        self._set_status("検索中...（%d件）" % len(parts))

        def do():
            if mode == "kishu":
                return core.search_kishu(parts, self.index)
            return core.search_chokujou(parts, self.index)
        self._run_async(do, on_done=self._show_results)

    def _show_results(self, rows):
        self.current_rows = rows
        mode = self.current_mode
        cols = core.KISHU_COLUMNS if mode == "kishu" else core.CHOKUJOU_COLUMNS
        keys = [k for k, _ in cols]

        # Treeview 列再構成
        self.tree.delete(*self.tree.get_children())
        self.tree["columns"] = keys
        widths = self._col_widths(mode)
        for k, title in cols:
            self.tree.heading(k, text=title)
            anchor = "e" if k in ("usage", "usage_total", "usage_direct") else "w"
            if k in ("dup",):
                anchor = "center"
            self.tree.column(k, width=widths.get(k, 110), anchor=anchor,
                             stretch=(k in ("part_name", "parent_name")))

        # 重複の連続値は空白化（Excel版の見やすさを踏襲）
        display = self._blank_duplicates(rows, mode, keys)
        cur_part = None
        band = 0
        for disp, raw in zip(display, rows):
            if raw["part"] != cur_part:
                cur_part = raw["part"]
                band ^= 1
            vals = [self._fmt(disp.get(k, "")) for k in keys]
            self.tree.insert("", tk.END, values=vals,
                             tags=("odd",) if band else ())

        n = len(rows)
        label = "機種コード逆展開検索" if mode == "kishu" else "直上親コード検索"
        self._set_status("%s 完了：%d行" % (label, n))
        if n == 0:
            messagebox.showinfo("結果", "該当する結果がありませんでした。")

    def _col_widths(self, mode):
        return {
            "part": 150, "model": 150, "usage": 90, "cucode": 150,
            "usage_total": 100, "parent": 150, "parent_name": 200,
            "usage_direct": 100, "part_name": 220, "parent_dk": 160,
            "dup": 50, "maker": 150,
        }

    def _fmt(self, v):
        if isinstance(v, float):
            return str(int(v)) if v == int(v) else ("%g" % v)
        return "" if v is None else str(v)

    def _blank_duplicates(self, rows, mode, keys):
        """連続する同一グループの先頭以外で重複セルを空白にした表示用コピーを返す。"""
        out = []
        prev = None
        for r in rows:
            d = dict(r)
            if prev is not None and r["part"] == prev["part"] and r["part"] != "":
                d["part"] = ""
                if "part_name" in d:
                    d["part_name"] = ""
                if mode == "chokujou":
                    same_model = r["model"] == prev["model"] and r["model"] != "該当なし"
                    if same_model:
                        d["model"] = ""
                        d["usage_total"] = ""
                        d["maker"] = ""
                        same_parent = (r["parent"] == prev["parent"]
                                       and r["parent"] not in ("", "該当なし"))
                        if same_parent:
                            d["parent"] = ""
                            d["parent_name"] = ""
                            d["usage_direct"] = ""
                            d["parent_dk"] = ""
                else:
                    if r["model"] == prev["model"]:
                        d["model"] = ""
            out.append(d)
            prev = r
        return out

    # --------------------------------------------------- 保存
    def on_save(self):
        if not self.current_rows:
            messagebox.showwarning(
                "確認", "保存する検索結果がありません。\n先に検索を実行してください。")
            return
        mode = self.current_mode
        cols = core.KISHU_COLUMNS if mode == "kishu" else core.CHOKUJOU_COLUMNS
        keys = [k for k, _ in cols]
        titles = [t for _, t in cols]
        default = ("逆検索結果_%s"
                   % datetime.datetime.now().strftime("%Y%m%d%H%M%S"))
        path = filedialog.asksaveasfilename(
            title="保存先とファイル名を指定してください",
            initialfile=default + ".xlsx",
            defaultextension=".xlsx",
            filetypes=[("Excel ブック", "*.xlsx"), ("CSV (Shift-JIS)", "*.csv")])
        if not path:
            return
        try:
            if path.lower().endswith(".csv"):
                self._save_csv(path, keys, titles)
            else:
                self._save_xlsx(path, keys, titles)
        except Exception as e:
            messagebox.showerror("エラー", "保存に失敗しました。\n%s" % e)
            return
        self._set_status("保存しました: %s" % path)
        messagebox.showinfo("完了", "別ファイルへの保存が完了しました。")

    def _save_csv(self, path, keys, titles):
        with open(path, "w", newline="", encoding="cp932", errors="replace") as f:
            w = csv.writer(f)
            w.writerow(titles)
            for r in self.current_rows:
                w.writerow([self._fmt(r.get(k, "")) for k in keys])

    def _save_xlsx(self, path, keys, titles):
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font, Alignment, PatternFill
            from openpyxl.utils import get_column_letter
        except ImportError:
            # openpyxl が無ければ CSV にフォールバック
            alt = os.path.splitext(path)[0] + ".csv"
            self._save_csv(alt, keys, titles)
            messagebox.showinfo(
                "Excel未対応", "openpyxl が無いため CSV で保存しました:\n%s" % alt)
            return
        wb = Workbook()
        ws = wb.active
        ws.title = "機種コード" if self.current_mode == "kishu" else "直上親コード"
        head_fill = PatternFill("solid", fgColor="2980B9")
        head_font = Font(color="FFFFFF", bold=True)
        ws.append(titles)
        for c in ws[1]:
            c.fill = head_fill
            c.font = head_font
            c.alignment = Alignment(horizontal="center")
        for r in self.current_rows:
            row = []
            for k in keys:
                v = r.get(k, "")
                if isinstance(v, float):
                    v = int(v) if v == int(v) else v
                row.append(v)
            ws.append(row)
        for i, k in enumerate(keys, start=1):
            ws.column_dimensions[get_column_letter(i)].width = \
                self._col_widths(self.current_mode).get(k, 110) / 7.0
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        wb.save(path)

    # --------------------------------------------------- その他
    def on_clear(self):
        self.txt.delete("1.0", tk.END)


def main():
    app = App()
    app.mainloop()


if __name__ == "__main__":
    main()
