# -*- coding: utf-8 -*-
"""
製品構成 逆検索ツール（ブラウザ版 / Flask + HTML）
  Excel/VBA版 製品構成逆検索ツール Ver1.04 をWebアプリへ移植したもの。
  検索ロジック(core.py)・マスタ取得層(dbsource.py)はデスクトップ版と共通。

  (1) 機種コード逆展開検索 : 部品 → 使用する最上位機種コード・使用数・客先品目コード
  (2) 直上親コード検索     : 部品 → 直上親(TD-仮親はスキップ)・名称・代表機種・メーカー
  (3) CSV / Excel ダウンロード

  マスタ取得は進捗付き（XITEM→XPRTS→XHEAD→XSECT→インデックス構築）。

DBへは SELECT のみ（書込みは行わない）。tdread は読取り専用ユーザー。
"""

import os
import io
import csv
import sys
import time
import datetime
import threading
import configparser

from flask import Flask, request, jsonify, render_template, Response

import core
import dbsource

# 通常実行と PyInstaller でのexe実行(frozen)の両対応
if getattr(sys, "frozen", False):
    EXE_DIR = os.path.dirname(sys.executable)              # exe の置き場所
    BUNDLE_DIR = getattr(sys, "_MEIPASS", EXE_DIR)         # 同梱ファイルの展開先
else:
    EXE_DIR = os.path.dirname(os.path.abspath(__file__))
    BUNDLE_DIR = EXE_DIR

# config.ini は「exeと同じ場所」を優先（利用者が編集可能）。無ければ同梱版を使う。
_cfg_path = os.path.join(EXE_DIR, "config.ini")
if not os.path.exists(_cfg_path):
    _cfg_path = os.path.join(BUNDLE_DIR, "config.ini")
cfg = configparser.ConfigParser()
cfg.read(_cfg_path, encoding="utf-8")

DEMO = ("--demo" in sys.argv) or (
    cfg.get("app", "demo", fallback="false").lower() == "true")

# テンプレートは同梱先（exe時は _MEIPASS/templates）から読む
app = Flask(__name__, template_folder=os.path.join(BUNDLE_DIR, "templates"))
# 画面(HTML/JS)が古いまま残らないようキャッシュを無効化
app.config["TEMPLATES_AUTO_RELOAD"] = True
app.config["SEND_FILE_MAX_AGE_DEFAULT"] = 0


@app.after_request
def _no_cache(resp):
    resp.headers["Cache-Control"] = "no-store, no-cache, must-revalidate, max-age=0"
    resp.headers["Pragma"] = "no-cache"
    return resp


# ---- マスタ（プロセス内に保持）と読込進捗 ----
_INDEX = {"obj": None}
_LOAD = {"state": "idle", "stage": "", "percent": 0, "summary": "",
         "error": "", "loaded_at": "", "elapsed": 0}
_LOCK = threading.Lock()

# 進捗メッセージ → パーセントの目安
_STAGE_PCT = [
    ("XITEM", 15), ("XPRTS", 45), ("XHEAD", 65),
    ("XSECT", 80), ("インデックス", 92),
]


def db_cfg():
    if not cfg.has_section("db"):
        return {}
    return dict(cfg.items("db"))


def _progress_cb(msg):
    pct = _LOAD["percent"]
    for key, p in _STAGE_PCT:
        if key in msg:
            pct = p
            break
    _LOAD["stage"] = msg
    _LOAD["percent"] = pct


def start_load():
    """マスタ読込をバックグラウンドで開始（多重起動しない）。"""
    with _LOCK:
        if _INDEX["obj"] is not None:
            _LOAD.update(state="done", percent=100)
            return
        if _LOAD["state"] == "loading":
            return
        _LOAD.update(state="loading", stage="接続を準備しています...",
                     percent=3, error="", summary="")

    def worker():
        t0 = time.time()
        try:
            if DEMO:
                _progress_cb("デモデータを準備中...")
                idx = dbsource.load_index_demo(progress=_progress_cb)
            else:
                idx = dbsource.load_index_from_db(db_cfg(), progress=_progress_cb)
            _INDEX["obj"] = idx
            n_item = len(idx.info_bunr)
            n_prts = sum(len(v) for v in idx.dict_fwd.values())
            elapsed = round(time.time() - t0, 1)
            now = datetime.datetime.now().strftime("%Y/%m/%d %H:%M:%S")
            _LOAD.update(state="done", stage="完了", percent=100,
                         summary="品目 %d件 / 構成 %d行" % (n_item, n_prts),
                         loaded_at=now, elapsed=elapsed)
        except Exception as e:
            _LOAD.update(state="error", stage="エラー", error=str(e))

    threading.Thread(target=worker, daemon=True).start()


def require_index():
    return _INDEX["obj"]


# ---------------- ルート ----------------
@app.route("/")
def index():
    return render_template("index.html", demo=DEMO)


@app.route("/api/start_load", methods=["POST"])
def api_start_load():
    start_load()
    return jsonify(_load_state())


@app.route("/api/load_progress")
def api_load_progress():
    return jsonify(_load_state())


def _load_state():
    d = dict(_LOAD)
    d["demo"] = DEMO
    d["loaded"] = _INDEX["obj"] is not None
    return d


@app.route("/api/status")
def status():
    return jsonify(_load_state())


def _run_search(mode, parts):
    idx = require_index()
    if idx is None:
        raise RuntimeError("マスタが読込まれていません。読込完了までお待ちください。")
    if mode == "kishu":
        return core.search_kishu(parts, idx), core.KISHU_COLUMNS
    return core.search_chokujou(parts, idx), core.CHOKUJOU_COLUMNS


def _parse_parts(payload):
    raw = payload.get("parts", "")
    lines = raw if isinstance(raw, list) else str(raw).splitlines()
    return [s.strip() for s in lines if s.strip()]


@app.route("/api/search", methods=["POST"])
def search():
    payload = request.get_json(force=True)
    mode = payload.get("mode", "kishu")
    parts = _parse_parts(payload)
    if not parts:
        return jsonify({"ok": False, "msg": "部品コードを入力してください。"}), 400
    if require_index() is None:
        return jsonify({"ok": False, "msg": "マスタを読込中です。完了までお待ちください。"}), 409
    try:
        rows, cols = _run_search(mode, parts)
    except Exception as e:
        return jsonify({"ok": False, "msg": "検索エラー: %s" % e}), 500
    return jsonify({
        "ok": True,
        "mode": mode,
        "columns": [{"key": k, "title": t} for k, t in cols],
        "rows": rows,
    })


@app.route("/api/download", methods=["POST"])
def download():
    payload = request.get_json(force=True)
    mode = payload.get("mode", "kishu")
    fmt = payload.get("format", "csv")
    cols_in = payload.get("columns")   # 画面が持っている列定義（再検索回避）
    rows_in = payload.get("rows")      # 画面が持っている検索結果（再検索回避）

    if isinstance(rows_in, list) and isinstance(cols_in, list) and cols_in:
        # 既に表示済みの結果をそのままファイル化（再検索しない＝高速）
        rows = rows_in
        keys = [c["key"] for c in cols_in]
        titles = [c["title"] for c in cols_in]
    else:
        # フォールバック：結果が渡されなければ検索して作る
        parts = _parse_parts(payload)
        if not parts:
            return jsonify({"ok": False, "msg": "部品コードを入力してください。"}), 400
        if require_index() is None:
            return jsonify({"ok": False, "msg": "マスタを読込中です。"}), 409
        rows, cols = _run_search(mode, parts)
        keys = [k for k, _ in cols]
        titles = [t for _, t in cols]
    base = "逆検索結果_" + ("機種コード" if mode == "kishu" else "直上親コード")

    if fmt == "xlsx":
        data = _build_xlsx(mode, rows, keys, titles)
        if data is not None:
            return Response(
                data,
                mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                headers={"Content-Disposition":
                         "attachment; filename=%s.xlsx" % base})

    buf = io.StringIO()
    w = csv.writer(buf, lineterminator="\r\n")
    w.writerow(titles)
    for r in rows:
        w.writerow([_fmt(r.get(k, "")) for k in keys])
    data = buf.getvalue().encode("cp932", errors="replace")
    return Response(data, mimetype="text/csv",
                    headers={"Content-Disposition":
                             "attachment; filename=%s.csv" % base})


def _fmt(v):
    if isinstance(v, float):
        return str(int(v)) if v == int(v) else ("%g" % v)
    return "" if v is None else str(v)


def _build_xlsx(mode, rows, keys, titles):
    """高速な write_only モードで .xlsx を生成（大量データでも速い）。"""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, Alignment, PatternFill
        from openpyxl.cell import WriteOnlyCell
    except ImportError:
        return None
    wb = Workbook(write_only=True)
    ws = wb.create_sheet("機種コード" if mode == "kishu" else "直上親コード")
    fill = PatternFill("solid", fgColor="2980B9")
    font = Font(color="FFFFFF", bold=True)
    align = Alignment(horizontal="center")
    header = []
    for t in titles:
        c = WriteOnlyCell(ws, value=t)
        c.fill = fill
        c.font = font
        c.alignment = align
        header.append(c)
    ws.append(header)
    for r in rows:
        row = []
        for k in keys:
            v = r.get(k, "")
            if isinstance(v, float):
                v = int(v) if v == int(v) else v
            row.append(v)
        ws.append(row)
    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue()


if __name__ == "__main__":
    port = int(cfg.get("app", "port", fallback="5001"))
    print("=" * 54)
    print(" 製品構成 逆検索ツール（ブラウザ版）")
    print(" モード:", "デモ(DB未接続)" if DEMO else "本番(SQL Server接続)")
    print(" ブラウザで http://127.0.0.1:%d を開いてください" % port)
    print("=" * 54)
    # 起動と同時にマスタ読込を先行開始（画面表示前から進捗を進める）
    start_load()
    # ブラウザを自動で開く（exe/通常実行とも）。サーバ起動直後に開くよう少し待つ。
    if "--no-browser" not in sys.argv:
        import webbrowser
        threading.Timer(
            1.2, lambda: webbrowser.open("http://127.0.0.1:%d" % port)).start()
    app.run(host="127.0.0.1", port=port, debug=False, threaded=True)
