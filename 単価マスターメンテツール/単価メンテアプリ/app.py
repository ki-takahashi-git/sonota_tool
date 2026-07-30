# -*- coding: utf-8 -*-
"""
単価マスター メンテナンスツール (Flask + HTML)
 (1) 閲覧ツール   : TPiCS SQL Server(XTANK/XHEAD/XITEM) を読取り専用で照会・警告表示
 (2) 一括登録CSV  : 改定単価を入力 → TPiCS テキスト読込用CSV(更新行+新規行)を生成
DBへは SELECT のみ。書込みは行わない（修正履歴を残すため TPiCS の読込で登録する）。
"""
import csv, io, os, sys, configparser
from flask import Flask, request, jsonify, render_template, Response
import core

BASE = os.path.dirname(os.path.abspath(__file__))
cfg = configparser.ConfigParser()
cfg.read(os.path.join(BASE, "config.ini"), encoding="utf-8")

LOGFILE = os.path.join(BASE, "app.log")
def _log(msg):
    """画面なし起動でも原因が追えるよう app.log に記録。"""
    import datetime
    line = "[%s] %s" % (datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S"), msg)
    try:
        with open(LOGFILE, "a", encoding="utf-8") as f:
            f.write(line + "\n")
    except Exception:
        pass
    try:
        print(line)
    except Exception:
        pass

DEMO = ("--demo" in sys.argv) or (cfg.get("app", "demo", fallback="false").lower() == "true")

app = Flask(__name__)
app.config["TEMPLATES_AUTO_RELOAD"] = True
app.jinja_env.auto_reload = True

@app.after_request
def _no_cache(resp):
    # 画面(HTML)はキャッシュさせず、修正がブラウザ更新だけで反映されるようにする
    if resp.mimetype == "text/html":
        resp.headers["Cache-Control"] = "no-store, must-revalidate"
        resp.headers["Pragma"] = "no-cache"
    return resp

# ---- 接続 ----
def get_conn():
    """pymssql優先、無ければpyodbcを試す。"""
    server = cfg.get("db", "server", fallback="YOUR_DB_SERVER_IP")
    port   = cfg.get("db", "port", fallback="1433")
    dbname = cfg.get("db", "database", fallback="TPICSDB_T")
    user   = cfg.get("db", "user", fallback="YOUR_DB_USER")
    pwd    = cfg.get("db", "password", fallback="YOUR_DB_PASSWORD")
    auth   = cfg.get("db", "auth", fallback="sql").lower()
    try:
        import pymssql
        kw = dict(server=server, port=port, database=dbname, timeout=60, login_timeout=15)
        if auth == "windows":
            return pymssql.connect(**kw)
        return pymssql.connect(user=user, password=pwd, **kw)
    except ImportError:
        import pyodbc
        drv = cfg.get("db", "odbc_driver", fallback="ODBC Driver 17 for SQL Server")
        if auth == "windows":
            cs = f"DRIVER={{{drv}}};SERVER={server},{port};DATABASE={dbname};Trusted_Connection=yes;"
        else:
            cs = f"DRIVER={{{drv}}};SERVER={server},{port};DATABASE={dbname};UID={user};PWD={pwd};"
        return pyodbc.connect(cs)

# 解決済み列名のキャッシュ
_SCHEMA = {}

def _columns_of(cur, table):
    cur.execute(
        "SELECT COLUMN_NAME FROM INFORMATION_SCHEMA.COLUMNS WHERE TABLE_NAME=%s" % _ph(),
        (table,))
    return [r[0] for r in cur.fetchall()]

def _ph():
    # pymssql は %s、pyodbc は ? のプレースホルダ
    try:
        import pymssql  # noqa
        return "%s"
    except ImportError:
        return "?"

def resolve_schema():
    """XTANK/XHEAD/XITEM の論理→実列名を解決（config上書き対応）。"""
    if _SCHEMA:
        return _SCHEMA
    if DEMO:
        _SCHEMA.update(demo_schema())
        return _SCHEMA
    conn = get_conn(); cur = conn.cursor()
    for key, tbl in (("tank","XTANK"),("head","XHEAD"),("item","XITEM"),("sect","XSECT")):
        cols = _columns_of(cur, tbl)
        res = core.resolve_columns(cols, core.CANDIDATES[key])
        # config.ini [columns] による上書き
        for logical in core.CANDIDATES[key]:
            ov = cfg.get("columns", f"{key}_{logical}", fallback=None)
            if ov:
                res[logical] = ov
        _SCHEMA[key] = {"table": tbl, "all": cols, "map": res}
    conn.close()
    return _SCHEMA

# ---------------- デモデータ ----------------
def demo_schema():
    base = ["CODE","VENDOR","TID","TDATE","EDATE","TVOL","TEMA","PRICE","SPRICE","TAXRATE","TAXKUBU","HIMOKU"]
    return {
        "tank": {"table":"XTANK","all":base,"map":{c:c for c in base}},
        "head": {"table":"XHEAD","all":["CODE","NAME","MAKERNO"],"map":{"CODE":"CODE","NAME":"NAME","MPN":"MAKERNO"}},
        "item": {"table":"XITEM","all":["CODE","BUMO","LOTS"],"map":{"CODE":"CODE","BUMO":"BUMO","LOTS":"LOTS"}},
        "sect": {"table":"XSECT","all":["BUMO","NAME"],"map":{"CODE":"BUMO","NAME":"NAME"}},
    }

DEMO_ROWS = [
    # 通常品（一致OK）
    dict(CODE="R-1005",MPN="MCR01MZPF1002",NAME="チップ抵抗1005",VENDOR="ROHM",VENDORNAME="ローム(株)",TID="500101",TDATE="202404011",EDATE="999999991",TVOL="4000",TEMA="0",PRICE="0.8",SPRICE="0.8",TAXRATE="10",TAXKUBU="1",HIMOKU="1",LOTS="4000"),
    # 鈑金（MOQ不一致：LOTS=500 だがTVOLに500なし）
    dict(CODE="SM-2210",MPN="YMD-BRK-A",NAME="ブラケットA 鈑金",VENDOR="V-YMD",VENDORNAME="ヤマダ製作所",TID="500210",TDATE="202401011",EDATE="999999991",TVOL="100",TEMA="0",PRICE="320",SPRICE="320",TAXRATE="10",TAXKUBU="1",HIMOKU="2",LOTS="500"),
    dict(CODE="SM-2210",MPN="YMD-BRK-A",NAME="ブラケットA 鈑金",VENDOR="V-YMD",VENDORNAME="ヤマダ製作所",TID="500211",TDATE="202401011",EDATE="999999991",TVOL="1000",TEMA="0",PRICE="290",SPRICE="290",TAXRATE="10",TAXKUBU="1",HIMOKU="2",LOTS="500"),
    # コネクタ（無期限が複数＝改定時に全件終了が必要）
    dict(CODE="CN-3050",MPN="5566-50A",NAME="コネクタ50P MOLEX",VENDOR="V-MLX",VENDORNAME="日本モレックス(同)",TID="500305",TDATE="202312011",EDATE="999999991",TVOL="100",TEMA="0",PRICE="45",SPRICE="45",TAXRATE="10",TAXKUBU="1",HIMOKU="1",LOTS="500"),
    dict(CODE="CN-3050",MPN="5566-50A",NAME="コネクタ50P MOLEX",VENDOR="V-MLX",VENDORNAME="日本モレックス(同)",TID="500306",TDATE="202312011",EDATE="999999991",TVOL="500",TEMA="0",PRICE="40",SPRICE="40",TAXRATE="10",TAXKUBU="1",HIMOKU="1",LOTS="500"),
    # 終息品（0円）
    dict(CODE="OB-9001",MPN="HRS-OLD-9001",NAME="旧基板ASSY(終息)",VENDOR="V-HRS",VENDORNAME="ヒロセ電機(株)",TID="500900",TDATE="202504011",EDATE="999999991",TVOL="1",TEMA="0",PRICE="0",SPRICE="0",TAXRATE="10",TAXKUBU="1",HIMOKU="1",LOTS="10"),
]

# ---------------- データ取得 ----------------
def fetch_rows(keyword="", codes=None, tokens=None):
    sch = resolve_schema()
    if DEMO:
        rows = DEMO_ROWS
        if codes:
            cs = set(str(c) for c in codes)
            rows = [r for r in rows if str(r["CODE"]) in cs]
            return _group(rows, sch)
        if tokens:
            ts = set(str(t) for t in tokens)
            rows = [r for r in rows if str(r["CODE"]) in ts or str(r.get("MPN","")) in ts]
            return _group(rows, sch)
        if keyword:
            k = keyword.lower()
            rows = [r for r in rows if k in r["CODE"].lower() or k in r.get("NAME","").lower() or k in r["VENDOR"].lower() or k in r.get("VENDORNAME","").lower() or k in r.get("MPN","").lower()]
        return _group(rows, sch)
    tm, hm, im = sch["tank"]["map"], sch["head"]["map"], sch["item"]["map"]
    sm = sch.get("sect", {}).get("map", {})
    sel = []
    sel.append(f"t.[{tm['CODE']}] AS CODE")
    sel.append(f"t.[{tm['VENDOR']}] AS VENDOR")
    for k in ("TID","TDATE","EDATE","TVOL","TEMA","SPRICE","PRICE","TAXRATE","TAXKUBU","HIMOKU"):
        if k in tm: sel.append(f"t.[{tm[k]}] AS {k}")
    if "NAME" in hm: sel.append(f"h.[{hm['NAME']}] AS NAME")
    # メーカー品目コード(MPN)：基本マスター → 詳細マスター → 注文コード(VCODE) の順で採用
    item_mpn = im.get("MPN")
    mpn_src = None
    if hm.get("MPN"):
        mpn_src = ("h", hm["MPN"])
    elif item_mpn:
        mpn_src = ("li", "makercd")   # 詳細マスターは集計サブクエリ経由
    elif tm.get("VCODE"):
        mpn_src = ("t", tm["VCODE"])
    if mpn_src:
        sel.append("li.makercd AS MPN" if mpn_src[0] == "li"
                   else f"{mpn_src[0]}.[{mpn_src[1]}] AS MPN")
    # 発注先名称（製造担当マスター XSECT を結合）
    sect_ok = bool(sm.get("CODE") and sm.get("NAME"))
    if sect_ok:
        sel.append("sx.vname AS VENDORNAME")
    # ロットサイズ（製造担当別に複数の可能性 → 集計）
    sel.append("li.maxlots AS LOTS, li.minlots AS LOTS_MIN, li.nlots AS LOTS_N")
    li_mpn = f", MAX([{item_mpn}]) AS makercd" if item_mpn else ""
    sx_join = ""
    if sect_ok:
        sx_join = (
            f"LEFT JOIN (SELECT [{sm['CODE']}] AS scode, MAX([{sm['NAME']}]) AS vname "
            f"FROM [{sch['sect']['table']}] WITH (NOLOCK) GROUP BY [{sm['CODE']}]) sx "
            f"ON sx.scode = t.[{tm['VENDOR']}]"
        )
    q = f"""
        SELECT TOP (2000) {', '.join(sel)}
        FROM [{sch['tank']['table']}] t WITH (NOLOCK)
        LEFT JOIN [{sch['head']['table']}] h WITH (NOLOCK) ON h.[{hm['CODE']}] = t.[{tm['CODE']}]
        LEFT JOIN (
            SELECT [{im['CODE']}] AS CODE, MAX([{im['LOTS']}]) AS maxlots,
                   MIN([{im['LOTS']}]) AS minlots, COUNT(DISTINCT [{im['LOTS']}]) AS nlots{li_mpn}
            FROM [{sch['item']['table']}] WITH (NOLOCK) GROUP BY [{im['CODE']}]
        ) li ON li.CODE = t.[{tm['CODE']}]
        {sx_join}
    """
    params = ()
    if codes:
        ph = ",".join([_ph()] * len(codes))
        q += f" WHERE t.[{tm['CODE']}] IN ({ph})"
        params = tuple(str(c) for c in codes)
    elif tokens:
        ph = ",".join([_ph()] * len(tokens))
        conds = [f"t.[{tm['CODE']}] IN ({ph})"]
        plist = [str(t) for t in tokens]
        if mpn_src:
            if mpn_src[0] == "li":
                conds.append(f"li.makercd IN ({ph})")
            else:
                conds.append(f"{mpn_src[0]}.[{mpn_src[1]}] IN ({ph})")
            plist += [str(t) for t in tokens]
        q += " WHERE " + " OR ".join(conds)
        params = tuple(plist)
    elif keyword:
        like = f"%{keyword}%"
        conds = [f"t.[{tm['CODE']}] LIKE {_ph()}", f"t.[{tm['VENDOR']}] LIKE {_ph()}"]
        if "NAME" in hm: conds.append(f"h.[{hm['NAME']}] LIKE {_ph()}")
        if mpn_src:
            conds.append((f"li.makercd LIKE {_ph()}") if mpn_src[0]=="li"
                         else f"{mpn_src[0]}.[{mpn_src[1]}] LIKE {_ph()}")
        if sect_ok:
            conds.append(f"sx.vname LIKE {_ph()}")
        q += " WHERE " + " OR ".join(conds)
        params = tuple([like]*len(conds))
    q += f" ORDER BY t.[{tm['CODE']}], t.[{tm['VENDOR']}], t.[{tm.get('TVOL','TVOL')}]"
    import time
    conn = get_conn(); cur = conn.cursor()
    try:
        cur.execute("SET TRANSACTION ISOLATION LEVEL READ UNCOMMITTED")
    except Exception as _e:
        print("[warn] isolation level:", _e)
    t0 = time.time()
    print("[search] keyword=%r 実行中..." % keyword)
    cur.execute(q, params)
    cols = [d[0] for d in cur.description]
    rows = [dict(zip(cols, r)) for r in cur.fetchall()]
    conn.close()
    print("[search] %d行 取得 (%.2f秒)" % (len(rows), time.time()-t0))
    return _group(rows, sch)

def _group(rows, sch):
    """(CODE,VENDOR) 単位にまとめ、各行へ表示用・警告を付与。"""
    groups = {}
    for r in rows:
        key = (str(r.get("CODE")), str(r.get("VENDOR")))
        groups.setdefault(key, []).append(r)
    out = []
    for (code, vendor), grp in groups.items():
        lots_info = {"max": grp[0].get("LOTS"),
                     "min": grp[0].get("LOTS_MIN", grp[0].get("LOTS")),
                     "n": grp[0].get("LOTS_N", 1) or 1}
        warns = core.compute_warnings(grp, lots_info)
        # 同じロット(適用数量)ごとにまとめ→適用開始→適用終了 の時系列で並べる
        def _dnum(v):
            d = "".join(ch for ch in str(v) if ch.isdigit())
            return int(d) if d else 0
        grp = sorted(grp, key=lambda r: (_f(r.get("TVOL")), _dnum(r.get("TDATE")), _dnum(r.get("EDATE"))))
        for r in grp:
            td = core.parse_date(r.get("TDATE"))
            ed = core.parse_date(r.get("EDATE"))
            item = {
                "CODE": code, "NAME": r.get("NAME",""), "MPN": r.get("MPN",""),
                "VENDOR": vendor, "VENDORNAME": r.get("VENDORNAME",""),
                "TID": r.get("TID",""),
                "TDATE": r.get("TDATE",""), "TDATE_disp": td["disp"], "TDATE_shift": td["shift"],
                "EDATE": r.get("EDATE",""), "EDATE_disp": ed["disp"], "EDATE_inf": ed["infinite"],
                "TVOL": core.clean_num(r.get("TVOL","")), "TEMA": r.get("TEMA",""),
                "PRICE": core.clean_num(r.get("PRICE","")), "SPRICE": core.clean_num(r.get("SPRICE","")),
                "TAXRATE": core.clean_num(r.get("TAXRATE","")), "TAXKUBU": r.get("TAXKUBU",""), "HIMOKU": r.get("HIMOKU",""),
                "LOTS": core.clean_num(r.get("LOTS","")),
                "moq_match": _moq_match(r.get("TVOL"), r.get("LOTS")),
                "is_zero": _is_zero(r.get("PRICE")),
                "group_warnings": warns,
            }
            out.append(item)
    return out

def _f(x):
    try:
        return float(str(x).replace(',', ''))
    except (TypeError, ValueError):
        return 0.0


def _moq_match(tvol, lots):
    try:
        return float(str(tvol).replace(",","")) == float(str(lots).replace(",",""))
    except (TypeError, ValueError):
        return None

def _is_zero(p):
    try:
        return float(str(p).replace(",","")) == 0
    except (TypeError, ValueError):
        return False

# ---------------- ルート ----------------
@app.route("/")
def index():
    return render_template("index.html", demo=DEMO)

@app.route("/api/status")
def status():
    info = {"demo": DEMO}
    if DEMO:
        sch = resolve_schema()
        info.update(ok=True, msg="デモモードで動作中（DB未接続）",
                    schema={k:{"table":v["table"],"map":v["map"],"all":v.get("all",[])} for k,v in sch.items()})
        return jsonify(info)
    try:
        sch = resolve_schema()
        info.update(ok=True, msg="DB接続OK",
                    schema={k:{"table":v["table"],"map":v["map"],"all":v.get("all",[])} for k,v in sch.items()})
    except Exception as e:
        info.update(ok=False, msg="DB接続エラー: %s" % e)
    return jsonify(info)

@app.route("/api/sample")
def sample():
    """指定マスターの先頭数行を返す（列名と中身の確認用・読取り専用）。"""
    key = request.args.get("key", "sect")
    if key not in ("tank", "head", "item", "sect"):
        return jsonify({"ok": False, "msg": "invalid key"}), 400
    try:
        sch = resolve_schema()
        tbl = sch[key]["table"]
        if DEMO:
            # デモ用ダミー（XSECT想定）
            if key == "sect":
                cols = ["BUMO", "NAME", "SCInOut"]
                rows = [["1000299", "（実DBで表示されます）", "0"]]
            else:
                cols = sch[key]["all"]; rows = []
            return jsonify({"ok": True, "table": tbl, "columns": cols, "rows": rows})
        conn = get_conn(); cur = conn.cursor()
        try:
            cur.execute("SET TRANSACTION ISOLATION LEVEL READ UNCOMMITTED")
        except Exception:
            pass
        cur.execute(f"SELECT TOP (5) * FROM [{tbl}] WITH (NOLOCK)")
        cols = [d[0] for d in cur.description]
        rows = [[("" if v is None else str(v)) for v in r] for r in cur.fetchall()]
        conn.close()
        return jsonify({"ok": True, "table": tbl, "columns": cols, "rows": rows})
    except Exception as e:
        return jsonify({"ok": False, "msg": str(e)}), 500


@app.route("/api/search_codes", methods=["POST"])
def search_codes():
    """アイテムコード または メーカー品目コードのリスト（貼り付け）で一括検索。"""
    payload = request.get_json(force=True)
    raw = payload.get("tokens", [])
    tokens = [str(t).strip() for t in raw if str(t).strip()]
    if not tokens:
        return jsonify({"ok": True, "rows": [], "tokens": []})
    try:
        rows = fetch_rows(tokens=tokens)
        found = set(str(r["CODE"]) for r in rows) | set(str(r.get("MPN","")) for r in rows)
        notfound = [t for t in tokens if t not in found]
        return jsonify({"ok": True, "rows": rows, "tokens": tokens, "notfound": notfound})
    except Exception as e:
        return jsonify({"ok": False, "msg": str(e)}), 500


@app.route("/api/search")
def search():
    kw = request.args.get("q","").strip()
    try:
        return jsonify({"ok": True, "rows": fetch_rows(kw)})
    except Exception as e:
        return jsonify({"ok": False, "msg": str(e)}), 500

@app.route("/api/generate", methods=["POST"])
def generate():
    """edits: [{current:{...}, kaitei_date, new_price, new_tvol, new_sprice, mode}]"""
    payload = request.get_json(force=True)
    edits = payload.get("edits", [])
    out_rows = []
    for e in edits:
        cur = e["current"]
        cur["TDATE_shift"] = core.parse_date(cur.get("TDATE")).get("shift","1")
        rows = core.make_rows(cur, e["kaitei_date"], e["new_price"],
                              new_tvol=e.get("new_tvol"), new_sprice=e.get("new_sprice"),
                              mode=e.get("mode","update"))
        out_rows.extend(rows)
    return jsonify({"ok": True, "rows": out_rows, "columns": core.OUTPUT_COLUMNS})

@app.route("/api/download_csv", methods=["POST"])
def download_csv():
    payload = request.get_json(force=True)
    edits = payload.get("edits", [])
    enc = payload.get("encoding","cp932")
    out_rows = []
    for e in edits:
        cur = e["current"]; cur["TDATE_shift"] = core.parse_date(cur.get("TDATE")).get("shift","1")
        out_rows.extend(core.make_rows(cur, e["kaitei_date"], e["new_price"],
                        new_tvol=e.get("new_tvol"), new_sprice=e.get("new_sprice"),
                        mode=e.get("mode","update")))
    buf = io.StringIO()
    w = csv.writer(buf, lineterminator="\r\n")
    w.writerow(core.OUTPUT_COLUMNS)
    for r in out_rows:
        w.writerow([r.get(c,"") for c in core.OUTPUT_COLUMNS])
    data = buf.getvalue().encode(enc, errors="replace")
    return Response(data, mimetype="text/csv",
        headers={"Content-Disposition": "attachment; filename=tanka_import.csv"})


# ---------------- Excel往復（出力／取込） ----------------
# 記入用Excelの列順：左に「記入欄」、右に「参照」「システム用（変更不可）」
INPUT_COLS = ["改定単価", "改定日(例 2026-07-01)", "新適用数量(空=現行)", "区分(価格改定/新規追加のみ)"]
REF_COLS   = ["アイテムコード", "品名", "メーカー品目", "発注先", "発注先名称",
              "適用開始", "適用終了", "適用数量", "現行発注単価", "現行標準単価"]
# make_rows に渡す機械項目（変更不可・取込で使用）
SYS_COLS   = ["TID", "_CODE", "_VENDOR", "_TDATE", "_EDATE", "_TVOL", "_TEMA",
              "_PRICE", "_SPRICE", "_TAXRATE", "_TAXKUBU", "_HIMOKU"]

def _item_to_row(it):
    return {
        "改定単価": "", "改定日(例 2026-07-01)": "", "新適用数量(空=現行)": "", "区分(価格改定/新規追加のみ)": "価格改定",
        "アイテムコード": it.get("CODE",""), "品名": it.get("NAME",""), "メーカー品目": it.get("MPN",""),
        "発注先": it.get("VENDOR",""), "発注先名称": it.get("VENDORNAME",""),
        "適用開始": it.get("TDATE_disp",""), "適用終了": it.get("EDATE_disp",""),
        "適用数量": it.get("TVOL",""), "現行発注単価": it.get("PRICE",""), "現行標準単価": it.get("SPRICE",""),
        "TID": it.get("TID",""), "_CODE": it.get("CODE",""), "_VENDOR": it.get("VENDOR",""),
        "_TDATE": it.get("TDATE",""), "_EDATE": it.get("EDATE",""), "_TVOL": it.get("TVOL",""),
        "_TEMA": it.get("TEMA",""), "_PRICE": it.get("PRICE",""), "_SPRICE": it.get("SPRICE",""),
        "_TAXRATE": it.get("TAXRATE",""), "_TAXKUBU": it.get("TAXKUBU",""), "_HIMOKU": it.get("HIMOKU",""),
    }

@app.route("/api/export_xlsx", methods=["POST"])
def export_xlsx():
    """選択した品目の現行 全TVOL行（TID付き）を記入用Excelに出力（取りこぼし防止）。"""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.worksheet.datavalidation import DataValidation
    except ImportError:
        return jsonify({"ok": False, "msg": "openpyxl が未導入です。初回セットアップを実行してください。"}), 500
    payload = request.get_json(force=True)
    sel_rows = payload.get("rows", [])
    pairs = set((str(r.get("CODE")), str(r.get("VENDOR"))) for r in sel_rows)
    codes = list({str(r.get("CODE")) for r in sel_rows})
    if not codes:
        return jsonify({"ok": False, "msg": "選択がありません。"}), 400
    allitems = fetch_rows(codes=codes)
    items = [it for it in allitems if (str(it["CODE"]), str(it["VENDOR"])) in pairs]
    items.sort(key=lambda x: (x["CODE"], x["VENDOR"], _f(x.get("TVOL"))))

    wb = Workbook(); ws = wb.active; ws.title = "改定入力"
    headers = INPUT_COLS + REF_COLS + SYS_COLS
    ws.append(headers)
    headfill = PatternFill("solid", fgColor="1F3864")
    infill   = PatternFill("solid", fgColor="FFF2CC")   # 記入欄=薄黄
    sysfill  = PatternFill("solid", fgColor="EDEDED")
    thin = Side(style="thin", color="CCCCCC"); border = Border(left=thin,right=thin,top=thin,bottom=thin)
    for ci, h in enumerate(headers, 1):
        c = ws.cell(row=1, column=ci); c.font = Font(bold=True, color="FFFFFF"); c.fill = headfill
        c.alignment = Alignment(horizontal="center", wrap_text=True); c.border = border
    for it in items:
        ws.append([_item_to_row(it).get(h, "") for h in headers])
    # 列幅・塗り・データ検証
    nrows = len(items)
    for ci, h in enumerate(headers, 1):
        letter = ws.cell(row=1, column=ci).column_letter
        ws.column_dimensions[letter].width = 16 if h in INPUT_COLS else (22 if h in ("品名","発注先名称") else 13)
        for ri in range(2, nrows + 2):
            cell = ws.cell(row=ri, column=ci); cell.border = border
            if h in INPUT_COLS: cell.fill = infill
            elif h in SYS_COLS: cell.fill = sysfill
    # 区分のドロップダウン
    if nrows:
        kubun_col = ws.cell(row=1, column=len(INPUT_COLS)).column_letter  # 4列目=区分
        dv = DataValidation(type="list", formula1='"価格改定,新規追加のみ"', allow_blank=True)
        ws.add_data_validation(dv); dv.add(f"{kubun_col}2:{kubun_col}{nrows+1}")
    ws.freeze_panes = "A2"
    # システム用列を折りたたみ（任意）
    bio = io.BytesIO(); wb.save(bio); bio.seek(0)
    return Response(bio.read(),
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=kaitei_input.xlsx"})

def _norm_date(v):
    """Excelの日付セル（datetime/文字列）を YYYY-MM-DD に正規化。"""
    import datetime
    if v in (None, ""):
        return ""
    if isinstance(v, (datetime.datetime, datetime.date)):
        return "%04d-%02d-%02d" % (v.year, v.month, v.day)
    sv = str(v).strip()
    digits = "".join(ch for ch in sv if ch.isdigit())
    if len(digits) >= 8:
        return "%s-%s-%s" % (digits[:4], digits[4:6], digits[6:8])
    return sv

def _edits_from_xlsx(file_storage):
    from openpyxl import load_workbook
    wb = load_workbook(file_storage, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return [], []
    header = [("" if h is None else str(h).strip()) for h in rows[0]]
    idx = {h: i for i, h in enumerate(header)}
    def g(r, name):
        i = idx.get(name)
        return ("" if i is None or i >= len(r) or r[i] is None else r[i])
    edits, skipped = [], 0
    for r in rows[1:]:
        if r is None or all(c is None for c in r):
            continue
        price = g(r, "改定単価")
        date  = _norm_date(g(r, "改定日(例 2026-07-01)"))
        if str(price).strip() == "" or date == "":
            skipped += 1
            continue
        kubun = str(g(r, "区分(価格改定/新規追加のみ)")).strip()
        mode = "new" if kubun.startswith("新規") else "update"
        tvol = g(r, "新適用数量(空=現行)")
        current = {
            "TID": g(r, "TID"), "CODE": g(r, "_CODE"), "VENDOR": g(r, "_VENDOR"),
            "TDATE": g(r, "_TDATE"), "EDATE": g(r, "_EDATE"), "TVOL": g(r, "_TVOL"),
            "TEMA": g(r, "_TEMA"), "PRICE": g(r, "_PRICE"), "SPRICE": g(r, "_SPRICE"),
            "TAXRATE": g(r, "_TAXRATE"), "TAXKUBU": g(r, "_TAXKUBU"), "HIMOKU": g(r, "_HIMOKU"),
        }
        edits.append({"current": {k: ("" if v is None else str(v)) for k, v in current.items()},
                      "kaitei_date": date, "new_price": str(price),
                      "new_tvol": ("" if str(tvol).strip()=="" else str(tvol)), "mode": mode})
    return edits, skipped

@app.route("/api/import_xlsx", methods=["POST"])
def import_xlsx():
    """記入済みExcelを取込み、更新行＋新規行を生成してプレビュー＋edits を返す。"""
    if "file" not in request.files:
        return jsonify({"ok": False, "msg": "ファイルがありません。"}), 400
    try:
        edits, skipped = _edits_from_xlsx(request.files["file"])
    except Exception as e:
        return jsonify({"ok": False, "msg": "Excelの読込に失敗: %s" % e}), 500
    out_rows = []
    for e in edits:
        cur = e["current"]; cur["TDATE_shift"] = core.parse_date(cur.get("TDATE")).get("shift","1")
        out_rows.extend(core.make_rows(cur, e["kaitei_date"], e["new_price"],
                        new_tvol=e.get("new_tvol"), new_sprice=e.get("new_sprice"),
                        mode=e.get("mode","update")))
    return jsonify({"ok": True, "edits": edits, "rows": out_rows,
                    "columns": core.OUTPUT_COLUMNS, "skipped": skipped})


def _open_browser_later(url, delay=1.5):
    import threading, webbrowser
    threading.Timer(delay, lambda: webbrowser.open(url)).start()

if __name__ == "__main__":
    host = cfg.get("app", "host", fallback="127.0.0.1")   # 全員で使う場合は 0.0.0.0
    port = int(cfg.get("app", "port", fallback="5000"))
    open_browser = cfg.get("app", "open_browser", fallback="true").lower() == "true"
    import socket
    try:
        lan_ip = socket.gethostbyname(socket.gethostname())
    except Exception:
        lan_ip = "127.0.0.1"
    local_url = "http://127.0.0.1:%d" % port
    print("=" * 60)
    print(" 単価マスター メンテナンスツール")
    print(" モード:", "デモ(DB未接続)" if DEMO else "本番(SQL Server接続)")
    print(" このPCから      :", local_url)
    if host == "0.0.0.0":
        print(" 同じ社内LANから :", "http://%s:%d" % (lan_ip, port))
        print("  （他PCは上のURLをお気に入り登録すれば設定不要で使えます）")
    print("=" * 60)
    _log("起動: mode=%s host=%s port=%d" % ("demo" if DEMO else "honban", host, port))
    if open_browser:
        _open_browser_later(local_url)
    try:
        try:
            from waitress import serve
            _log("サーバー: waitress（複数人対応）")
            serve(app, host=host, port=port, threads=8)
        except ImportError:
            _log("サーバー: Flask標準（waitress未導入）")
            app.run(host=host, port=port, debug=False)
    except Exception as e:
        import traceback
        _log("致命的エラーで起動できませんでした:\n" + traceback.format_exc())
        raise
